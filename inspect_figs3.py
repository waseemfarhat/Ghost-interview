import docx
from docx.shared import Emu

DOC = r'Ghost_Interview_FYP_Evaluation3_Documentation.docx'
d = docx.Document(DOC)

# Build a linear list of body-level paragraphs
paras = d.paragraphs

def para_text(p):
    return p.text.strip()

# For each inline shape, find its caption: the nearest following paragraph
# that starts with "Figure" OR the immediately next non-empty paragraph.
from docx.oxml.ns import qn

# Map run element -> index among d.paragraphs for locating following caption
# Simpler: iterate body children in order
body = d.element.body

# collect ordered elements: paragraphs and inline shapes live inside runs
order = []
for child in body.iterchildren():
    tag = child.tag
    if tag == qn('w:p'):
        # find text
        texts = child.findall('.//' + qn('w:t'))
        txt = ''.join(t.text or '' for t in texts)
        has_pic = child.find('.//' + qn('w:drawing')) is not None
        order.append(('p', txt.strip(), has_pic))
    elif tag == qn('w:tbl'):
        order.append(('tbl', '<TABLE>', False))

# Now pair each picture paragraph with the next caption paragraph
figs = []
i = 0
while i < len(order):
    kind, txt, has_pic = order[i]
    if has_pic:
        # look ahead for caption
        cap = ''
        for j in range(i+1, min(i+4, len(order))):
            k2, t2, hp2 = order[j]
            if t2.lower().startswith('figure'):
                cap = t2
                break
            if k2 == 'p' and t2:
                cap = t2
                break
        figs.append(cap)
    i += 1

print('TOTAL FIGURE SLOTS:', len(figs))
for idx, c in enumerate(figs, 1):
    print(f'{idx:02d}: {c}')
