import docx, zipfile, io, os
from PIL import Image
DOC="Ghost_Interview_FYP_Evaluation3_Documentation.docx"
d=docx.Document(DOC)
print("PARAS", len(d.paragraphs), "TABLES", len(d.tables))
z=zipfile.ZipFile(DOC)
media=[n for n in z.namelist() if n.startswith("word/media/")]
ph=0; real=0
for n in media:
    im=Image.open(io.BytesIO(z.read(n)))
    if im.size==(1274,798): ph+=1
    else: real+=1
print("MEDIA_FILES", len(media), "REAL", real, "PLACEHOLDER_FILES", ph)
print("SIZE_MB", round(os.path.getsize(DOC)/1048576,2))
