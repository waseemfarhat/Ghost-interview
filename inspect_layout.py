import docx
from docx.oxml.ns import qn

DOC = 'Ghost_Interview_FYP_Evaluation3_Documentation.docx'
d = docx.Document(DOC)

def has_img(p):
    return bool(p._p.findall('.//' + qn('w:drawing')))

body = d.element.body
idx = 0
for p in d.paragraphs:
    img = has_img(p)
    txt = p.text.strip()
    style = p.style.name if p.style else ''
    marker = ' [IMAGE]' if img else ''
    if img or txt.startswith('Figure') or txt.lower().startswith('chapter') or txt.startswith('2.') or txt.startswith('3.'):
        print(f"{idx:3} | {style:14} | {txt[:90]!r}{marker}")
    idx += 1

print("\nTOTAL PARAGRAPHS:", idx)
print("TOTAL TABLES:", len(d.tables))
print("TOTAL IMAGES(inline):", sum(1 for p in d.paragraphs if has_img(p)))
