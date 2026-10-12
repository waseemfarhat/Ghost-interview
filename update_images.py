import docx, os, shutil
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image

DOC = 'Ghost_Interview_FYP_Evaluation3_Documentation.docx'
IMG = 'extracted_images'
MAX_W, MAX_H = 15.5, 12.5  # cm

def fit_size(path):
    w_px, h_px = Image.open(path).size
    ratio = w_px / h_px
    w = MAX_W; h = w / ratio
    if h > MAX_H:
        h = MAX_H; w = h * ratio
    return Cm(w), Cm(h)

def sync_extent(inline, cx, cy):
    inline.width = cx; inline.height = cy
    pic = inline._inline.graphic.graphicData.pic
    xfrm = pic.spPr.find(qn('a:xfrm'))
    if xfrm is not None:
        ext = xfrm.find(qn('a:ext'))
        if ext is not None:
            ext.set('cx', str(int(cx))); ext.set('cy', str(int(cy)))

# backup
shutil.copy2(DOC, DOC.replace('.docx', '_backup.docx'))
doc = docx.Document(DOC)

# ---------- 1. Replace 4 real UI mockups ----------
REPLACE = {
    'Figure 2.10:': 'scope_p11_14.jpg',   # Job-role selection screen
    'Figure 2.12:': 'scope_p11_15.jpg',   # Live interview session (avatar + webcam)
    'Figure 2.15:': 'scope_p12_16.jpg',   # Post-interview emotion & speech analytics
    'Figure 2.19:': 'scope_p12_17.jpg',   # Final confidence & competency report
}
paras = doc.paragraphs
shapes = doc.inline_shapes
shape_idx = 0; replaced = 0
for i, p in enumerate(paras):
    if not p._p.findall('.//' + qn('w:drawing')):
        continue
    cap = ''
    for j in range(i + 1, min(i + 4, len(paras))):
        if paras[j].text.strip():
            cap = paras[j].text.strip(); break
    for key, fn in REPLACE.items():
        if cap.startswith(key):
            path = os.path.join(IMG, fn)
            rId, _ = doc.part.get_or_add_image(path)
            inline = shapes[shape_idx]
            blip = inline._inline.graphic.graphicData.pic.blipFill.blip
            blip.set(qn('r:embed'), rId)
            w, h = fit_size(path)
            sync_extent(inline, w, h)
            replaced += 1
            print(f'  [OK] {key} -> {fn}  ({os.path.getsize(path)//1024} KB)')
            break
    shape_idx += 1
print(f'Replacements: {replaced}/4   (total image slots: {shape_idx})')

# ---------- 2. Appendix A: design diagrams ----------
GREEN = RGBColor(0x14, 0x53, 0x2D)
GREY = RGBColor(0x55, 0x55, 0x55)
DARK = RGBColor(0x33, 0x33, 0x33)

def page_break():
    r = doc.add_paragraph().add_run()
    br = OxmlElement('w:br'); br.set(qn('w:type'), 'page'); r._r.append(br)

def h1(text):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text); r.bold = True; r.font.size = Pt(16)
    r.font.name = 'Calibri'; r.font.color.rgb = GREEN

def sub(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text); r.bold = True; r.font.size = Pt(11.5)
    r.font.name = 'Calibri'; r.font.color.rgb = DARK

def figure(fn, cap):
    path = os.path.join(IMG, fn)
    w, h = fit_size(path)
    pp = doc.add_paragraph(); pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pp.paragraph_format.space_after = Pt(2)
    pp.add_run().add_picture(path, width=w, height=h)
    cp = doc.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(12)
    cr = cp.add_run(cap); cr.italic = True; cr.font.size = Pt(9)
    cr.font.name = 'Calibri'; cr.font.color.rgb = GREY

DIAGRAMS = [
    ('System Architecture', 'sdd1_p06_03.jpg',
     'System architecture of the Ghost Interview platform (presentation, application, AI-service and data layers).'),
    ('Context Diagram (Level-0 DFD)', 'sdd1_p05_02.jpg',
     'Context (Level-0) data-flow diagram showing the system boundary and external entities.'),
    ('Class Diagram', 'sdd1_p12_07.jpg',
     'Class diagram of the core domain model: User, Admin, InterviewSession, JobRole, SpeechAnalyzer, EmotionAnalyzer and Report.'),
    ('Use Case Diagram', 'srs_p06_02.jpg',
     'Use-case diagram showing candidate and administrator interactions with the system.'),
    ('User Management Module Flowchart', 'sdd1_p08_04.jpg',
     'Activity flowchart of the User Management module (registration and login flows).'),
    ('AI Interview Core Module Flowchart', 'sdd1_p10_05.jpg',
     'Activity flowchart of the AI Interview Core module (question generation, speech, emotion and scoring).'),
    ('Admin Module Flowchart', 'sdd1_p11_06.jpg',
     'Activity flowchart of the Admin module (user, job-role, analytics and reporting actions).'),
    ('Login Sequence Diagram', 'sdd1_p13_08.jpg',
     'Sequence diagram for the candidate login and session-token generation flow.'),
    ('Interview Session Sequence Diagram', 'sdd1_p14_09.jpg',
     'Sequence diagram for a complete interview session across frontend, backend, AI services and database.'),
    ('Project Timeline (Gantt Chart)', 'scope_p10_13.jpg',
     'Project Gantt chart showing the development phases and milestones.'),
]

page_break()
h1('Appendix A: System Design & Architecture Diagrams')
intro = doc.add_paragraph()
ir = intro.add_run('The diagrams below present the design artifacts that underpin the '
                   'implemented modules described in Chapter 2.')
ir.font.size = Pt(10.5); ir.font.name = 'Calibri'; ir.font.color.rgb = DARK
for title, fn, cap in DIAGRAMS:
    sub(title); figure(fn, cap)

doc.save(DOC)
print('SAVED:', DOC)
print('Paragraphs:', len(doc.paragraphs), ' Inline images:', len(doc.inline_shapes))
