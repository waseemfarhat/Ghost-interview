import docx
from docx.oxml.ns import qn

MINE = r'Ghost_Interview_FYP_Evaluation3_Documentation.docx'
d = docx.Document(MINE)

def fmt(p):
    s = set()
    szs = set()
    col = None
    for r in p.runs:
        if r.bold: s.add('B')
        if r.italic: s.add('I')
        if r.font.size: szs.add(r.font.size.pt)
        if r.font.color and r.font.color.rgb: col = str(r.font.color.rgb)
    a = p.alignment
    return f"[{'/'.join(sorted(s)) if s else '-'}|{sorted(szs)}|{col}|{a}]"

print("PARAS:", len(d.paragraphs), "TABLES:", len(d.tables), "IMGS:", len(d.inline_shapes))
print("===== PARAGRAPHS (non-empty) =====")
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    has_img = p._p.findall('.//' + qn('a:blip'))
    if not t and not has_img:
        continue
    tag = ' <IMG>' if has_img else ''
    print(f"{i:03d} {fmt(p)} {t[:95]}{tag}")
